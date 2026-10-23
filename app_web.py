import io
import csv
import random
from datetime import datetime, date, timedelta
from decimal import Decimal, ROUND_HALF_UP, ROUND_FLOOR, ROUND_CEILING
import streamlit as st

# Wordライブラリ
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ページ基本設定
st.set_page_config(
    page_title="計算実務2級 問題ジェネレーター",
    page_icon="📝",
    layout="wide"
)

# ---------------------------------------------------------
# 1. 多言語翻訳辞書（全9言語対応）
# ---------------------------------------------------------
LANG_OPTIONS = [
    "なし（日本語のみ）",
    "英語 (English)",
    "ネパール語 (नेपाली)",
    "シンハラ語 (සිංහල)",
    "ベトナム語 (Tiếng Việt)",
    "中国語 (中文)",
    "ミャンマー語 (မြန်မာ)",
    "タイ語 (ไทย)",
    "タガログ語 (Tagalog)",
    "ベンガル語 (বাংলা)"
]

TRANSLATION_DICT = {
    "英語 (English)": {
        "（円未満切り上げ）": " (rounded up to nearest yen)",
        "（円未満切り捨て）": " (rounded down to nearest yen)",
        "（円未満四捨五入）": " (rounded to nearest yen)",
        "（セント未満四捨五入）": " (rounded to nearest cent)",
        "（￥100未満切り上げ）": " (rounded up to nearest 100 yen)",
        "（片落とし）": " (excluding first day)",
        "（両端入れ）": " (inclusive of both days)",
        "原価をXとすると": "Let X be the cost",
        "定価をXとすると": "Let X be the list price",
        "売上高をXとすると": "Let X be the sales amount",
        "元金をXとすると": "Let X be the principal",
        "手形金額をXとすると": "Let X be the face value",
        "ダース": " doz", "個": " pcs", "か月": " months", "年": " years ",
        "日後": " days later", "日前": " days before", "日間": " days", "日": " days",
        "純量": "Net weight ", "利息": "Interest ", "定価": "List price ",
        "原価": "Cost ", "売上高": "Sales ", "差額": "Difference ",
        "値引時1kgあたり": "Discount price/kg: ", "増量時1kgあたり": "Extra volume price/kg: ",
        "値引有利": "Discount is more profitable", "増量有利": "Extra volume is more profitable"
    },
    "ネパール語 (नेपाली)": {
        "（円未満切り上げ）": " (येन भन्दा कम राउन्ड अप)",
        "（円未満切り捨て）": " (येन भन्दा कम हटाइयो)",
        "（円未満四捨五入）": " (निकटतम येनमा गोलो गरियो)",
        "（セント未満四捨五入）": " (निकटतम सेन्टमा गोलो गरियो)",
        "（￥100未満切り上げ）": " (१०० येन भन्दा कम राउन्ड अप)",
        "（片落とし）": " (पहिलो दिन बाहेक)",
        "（両端入れ）": " (दुबै दिन समावेश)",
        "原価をXとすると": "लागतलाई X मान्दा",
        "定価をXとすると": "सूची मूल्यलाई X मान्दा",
        "売上高をXとすると": "बिक्री रकमलाई X मान्दा",
        "元金をXとすると": "साँवा रकमलाई X मान्दा",
        "手形金額をXとすると": "बिल रकमलाई X मान्दा",
        "ダース": " दर्जन", "個": " वटा", "か月": " महिना", "年": " वर्ष ",
        "日後": " दिन पछि", "日前": " दिन अघि", "日間": " दिन", "日": " दिन",
        "純量": "खुद वजन ", "利息": "ब्याज ", "定価": "सूची मूल्य ",
        "原価": "लागत ", "売上高": "बिक्री ", "差額": "फरक ",
        "値引時1kgあたり": "छुट मूल्य/केजी: ", "増量時1kgあたり": "थप परिमाण मूल्य/केजी: ",
        "値引有利": "छुट बढी लाभदायक छ", "増量有利": "थप परिमाण बढी लाभदायक छ"
    },
    "シンハラ語 (සිංහල)": {
        "（円未満切り上げ）": " (යෙන් අගයට වැඩි කිරීම)",
        "（円未満切り捨て）": " (යෙන් අගයට අඩු කිරීම)",
        "（円未満四捨五入）": " (ආසන්නතම යෙන් අගයට)",
        "（セント未満四捨五入）": " (ආසන්නතම ශත අගයට)",
        "（￥100未満切り上げ）": " (යෙන් 100 දක්වා වැඩි කිරීම)",
        "（片落とし）": " (පළමු දිනය හැර)",
        "（両端入れ）": " (දින දෙකම ඇතුළත්ව)",
        "原価をXとすると": "පිරිවැය X ලෙස ගනිමු",
        "定価をXとすると": "නියමිත මිල X ලෙස ගනිමු",
        "売上高をXとすると": "විකුණුම් මුදල X ලෙස ගනිමු",
        "元金をXとすると": "මූලධනය X ලෙස ගනිමු",
        "手形金額をXとすると": "බිල්පත් මුදල X ලෙස ගනිමු",
        "ダース": " ඩසන්", "個": " ප්‍රමාණ", "か月": " මාස", "年": " වසර ",
        "日後": " දිනකට පසු", "日前": " දිනකට පෙර", "日間": " දින", "日": " දින",
        "純量": "ශුද්ධ බර ", "利息": "පොලිය ", "定価": "නියමිත මිල ",
        "原価": "පිරිවැය ", "売上高": "විකිණුම් ", "差額": "වෙනස ",
        "値引時1kgあたり": "වට්ටම් මිල/කි.ග්‍රෑ: ", "増量時1kgあたり": "වැඩි ප්‍රමාණ මිල/කි.ග්‍රෑ: ",
        "値引有利": "වට්ටම් වාසිදායක වේ", "増量有利": "වැඩි ප්‍රමාණය වාසිදායක වේ"
    },
    "ベトナム語 (Tiếng Việt)": {
        "（円未満切り上げ）": " (làm tròn lên hàng yên)",
        "（円未満切り捨て）": " (làm tròn xuống hàng yên)",
        "（円未満四捨五入）": " (làm tròn đến hàng yên gần nhất)",
        "（セント未満四捨五入）": " (làm tròn đến cent gần nhất)",
        "（￥100未満切り上げ）": " (làm tròn lên 100 yên)",
        "（片落とし）": " (không tính ngày đầu)",
        "（両端入れ）": " (tính cả hai ngày)",
        "原価をXとすると": "Gọi giá vốn là X",
        "定価をXとすると": "Gọi giá niêm yết là X",
        "売上高をXとすると": "Gọi doanh số là X",
        "元金をXとすると": "Gọi tiền gốc là X",
        "手形金額をXとすると": "Gọi mệnh giá hối phiếu là X",
        "ダース": " tá", "個": " cái", "か月": " tháng", "年": " năm ",
        "日後": " ngày sau", "日前": " ngày trước", "日間": " ngày", "日": " ngày",
        "純量": "Trọng lượng tịnh ", "利息": "Tiền lãi ", "定価": "Giá niêm yết ",
        "原価": "Giá vốn ", "売上高": "Doanh số ", "差額": "Chênh lệch ",
        "値引時1kgあたり": "Giá giảm/1kg: ", "増量時1kgあたり": "Giá tăng lượng/1kg: ",
        "値引有利": "Giảm giá có lợi hơn", "増量有利": "Tăng số lượng có lợi hơn"
    },
    "中国語 (中文)": {
        "（円未満切り上げ）": "（不足1日元进位）",
        "（円未満切り捨て）": "（不足1日元舍去）",
        "（円未満四捨五入）": "（四舍五入到日元）",
        "（セント未満四捨五入）": "（四舍五入到美分）",
        "（￥100未満切り上げ）": "（不足100日元进位）",
        "（片落とし）": "（算尾不算头）",
        "（両端入れ）": "（两头都算）",
        "原価をXとすると": "设进价为X",
        "定価をXとすると": "设标价为X",
        "売上高をXとすると": "设销售额为X",
        "元金をXとすると": "设本金为X",
        "手形金額をXとすると": "设票据金额为X",
        "ダース": "打", "個": "个", "か月": "个月", "年": "年",
        "日後": "天后", "日前": "天前", "日間": "天", "日": "天",
        "純量": "净重", "利息": "利息", "定価": "标价",
        "原価": "进价", "売上高": "销售额", "差額": "差额",
        "値引時1kgあたり": "打折后每公斤: ", "増量時1kgあたり": "增量后每公斤: ",
        "値引有利": "打折更有利", "増量有利": "增量更有利"
    },
    "ミャンマー語 (မြန်မာ)": {
        "（円未満切り上げ）": " (ယန်းအောက် အပေါ်တင်)",
        "（円未満切り捨て）": " (ယန်းအောက် ဖြုတ်ပစ်)",
        "（円未満四捨五入）": " (အနီးစပ်ဆုံး ယန်းသို့ ပေါင်းတင်)",
        "（セント未満四捨五入）": " (အနီးစပ်ဆုံး ဆင့်သို့ ပေါင်းတင်)",
        "（￥100未満切り上げ）": " (ယန်း ၁၀၀ အောက် အပေါ်တင်)",
        "（片落とし）": " (ပထမရက် မပါ)",
        "（両端入れ）": " (နှစ်ရက်လုံး အပါအဝင်)",
        "原価をXとすると": "အရင်းအနှီးကို X ဟုထားပါ",
        "定価をXとすると": "သတ်မှတ်ဈေးကို X ဟုထားပါ",
        "売上高をXとすると": "အရောင်းပမာဏကို X ဟုထားပါ",
        "元金をXとすると": "အရင်းငွေကို X ဟုထားပါ",
        "手形金額をXとすると": "ငွေတောင်းခံလွှာပမာဏကို X ဟုထားပါ",
        "ダース": " ဒါဇင်", "個": " ခု", "か月": " လ", "年": " နှစ် ",
        "日後": " ရက်အကြာ", "日前": " ရက်အလို", "日間": " ရက်", "日": " ရက်",
        "純量": "အသားတင်အလေးချိန် ", "利息": "အတိုး ", "定価": "သတ်မှတ်ဈေး ",
        "原価": "အရင်း ", "売上高": "အရောင်း ", "差額": "ကွာခြားချက် ",
        "値引時1kgあたり": "လျှော့ဈေး ၁ ကီလို: ", "増量時1kgあたり": "အပိုပမာဏ ၁ ကီလို: ",
        "値引有利": "ဈေးလျှော့ခြင်းက ပိုအကျိုးရှိသည်", "増量有利": "အလေးချိန်တိုးခြင်းက ပိုအကျိုးရှိသည်"
    },
    "タイ語 (ไทย)": {
        "（円未満切り上げ）": " (ปัดขึ้นเป็นจำนวนเต็มเยน)",
        "（円未満切り捨て）": " (ปัดเศษเยนทิ้ง)",
        "（円未満四捨五入）": " (ปัดเศษเป็นเยนที่ใกล้ที่สุด)",
        "（セント未満四捨五入）": " (ปัดเศษเป็นเซนต์ที่ใกล้ที่สุด)",
        "（￥100未満切り上げ）": " (ปัดขึ้นเป็นจำนวนเต็ม 100 เยน)",
        "（片落とし）": " (ไม่นับวันแรก)",
        "（両端入れ）": " (นับรวมทั้งวันแรกและวันสุดท้าย)",
        "原価をXとすると": "กำหนดให้ต้นทุนคือ X",
        "定価をXとすると": "กำหนดให้ราคาป้ายคือ X",
        "売上高をXとすると": "กำหนดให้ยอดขายคือ X",
        "元金をXとすると": "กำหนดให้เงินต้นคือ X",
        "手形金額をXとすると": "กำหนดให้มูลค่าตั๋วเงินคือ X",
        "ダース": " โหล", "個": " ชิ้น", "か月": " เดือน", "年": " ปี ",
        "日後": " วันหลังจากนี้", "日前": " วันก่อน", "日間": " วัน", "日": " วัน",
        "純量": "น้ำหนักสุทธิ ", "利息": "ดอกเบี้ย ", "定価": "ราคาป้าย ",
        "原価": "ต้นทุน ", "売上高": "ยอดขาย ", "差額": "ส่วนต่าง ",
        "値引時1kgあたり": "ราคาลด/กก.: ", "増量時1kgあたり": "ราคาเพิ่มปริมาณ/กก.: ",
        "値引有利": "การลดราคามีกำไรมากกว่า", "増量有利": "การเพิ่มปริมาณมีกำไรมากกว่า"
    },
    "タガログ語 (Tagalog)": {
        "（円未満切り上げ）": " (i-round up sa pinakamalapit na yen)",
        "（円未満切り捨て）": " (i-round down sa pinakamalapit na yen)",
        "（円未満四捨五入）": " (i-round off sa pinakamalapit na yen)",
        "（セント未満四捨五入）": " (i-round off sa pinakamalapit na sentimo)",
        "（￥100未満切り上げ）": " (i-round up sa pinakamalapit na 100 yen)",
        "（片落とし）": " (hindi kasama ang unang araw)",
        "（両端入れ）": " (kasama ang parehong araw)",
        "原価をXとすると": "Kung ang puhunan ay X",
        "定価をXとすると": "Kung ang itinakdang presyo ay X",
        "売上高をXとすると": "Kung ang kabuuang benta ay X",
        "元金をXとすると": "Kung ang prinsipal ay X",
        "手形金額をXとすると": "Kung ang halaga ng promissory note ay X",
        "ダース": " dosena", "個": " piraso", "か月": " buwan", "年": " taon ",
        "日後": " araw pagkatapos", "日前": " araw bago", "日間": " araw", "日": " araw",
        "純量": "Netong timbang ", "利息": "Interes ", "定価": "Itinakdang Presyo ",
        "原価": "Puhunan ", "売上高": "Benta ", "差額": "Pagkakaiba ",
        "値引時1kgあたり": "Presyo bawat kg (Diskwento): ", "増量時1kgあたり": "Presyo bawat kg (Dagdag-dami): ",
        "値引有利": "Mas kapaki-pakinabang ang Diskwento", "増量有利": "Mas kapaki-pakinabang ang Dagdag-dami"
    },
    "ベンガル語 (বাংলা)": {
        "（円未満切り上げ）": " (ইয়েন রাউন্ড আপ)",
        "（円未満切り捨て）": " (ইয়েন রাউন্ড ডাউন)",
        "（円未満四捨五入）": " (নিকটতম ইয়েনে রাউন্ড অফ)",
        "（セント未満四捨五入）": " (নিকটতম সেন্টে রাউন্ড অফ)",
        "（￥100未満切り上げ）": " (১০০ ইয়েন রাউন্ড আপ)",
        "（片落とし）": " (প্রথম দিন বাদ দিয়ে)",
        "（両端入れ）": " (উভয় দিন অন্তর্ভুক্ত)",
        "原価をXとすると": "ক্রয়মূল্য X ধরা হলে",
        "定価をXとすると": "ধার্যমূল্য X ধরা হলে",
        "売上高をXとすると": "বিক্রয়মূল্য X ধরা হলে",
        "元金をXとすると": "আসল টাকা X ধরা হলে",
        "手形金額をXとすると": "বিল মূল্য X ধরা হলে",
        "ダース": " ডজন", "個": " টি", "か月": " মাস", "年": " বছর ",
        "日後": " দিন পর", "日前": " দিন আগে", "日間": " দিন", "日": " দিন",
        "純量": "প্রকৃত ওজন ", "利息": "সুদ ", "定価": "ধার্যমূল্য ",
        "原価": "ক্রয়মূল্য ", "売上高": "বিক্রয় ", "差額": "পার্থক্য ",
        "値引時1kgあたり": "ছাড়ের পর প্রতি কেজি: ", "増量時1kgあたり": "পরিমাণ বৃদ্ধির পর প্রতি কেজি: ",
        "値引有利": "মূল্যছাড় বেশি লাভজনক", "増量有利": "পরিমাণ বৃদ্ধি বেশি লাভজনক"
    }
}

def translate_explanation(jp_text: str, target_lang: str) -> str:
    if target_lang not in TRANSLATION_DICT:
        return jp_text
    res = jp_text
    for k, v in TRANSLATION_DICT[target_lang].items():
        res = res.replace(k, v)
    return res.replace("×", " * ").replace("÷", " / ").replace("＝", " = ").replace("－", " - ").replace("＋", " + ").replace("→", " -> ")

# ---------------------------------------------------------
# 2. 計算・日付補助関数
# ---------------------------------------------------------
def round_yen(val: Decimal, mode: str) -> int:
    if mode == "四捨五入":
        return int(val.quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    elif mode == "切り捨て":
        return int(val.quantize(Decimal('1'), rounding=ROUND_FLOOR))
    elif mode == "切り上げ":
        return int(val.quantize(Decimal('1'), rounding=ROUND_CEILING))
    raise ValueError

def round_cent(val: Decimal) -> Decimal:
    return val.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def round_up_100yen(val: Decimal) -> int:
    q = (val / Decimal('100')).quantize(Decimal('1'), rounding=ROUND_CEILING)
    return int(q) * 100

def get_date_pair(start_month: int, start_day: int, days: int, is_ryotan: bool):
    d1 = date(2025, start_month, start_day)
    delta = days - 1 if is_ryotan else days
    d2 = d1 + timedelta(days=delta)
    month_days = [0, 31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    breakdown = []
    for m in range(d1.month, d2.month + 1):
        if m == d1.month and m == d2.month:
            breakdown.append(str(days))
        elif m == d1.month:
            rem = month_days[m] - d1.day + (1 if is_ryotan else 0)
            breakdown.append(str(rem))
        elif m == d2.month:
            breakdown.append(str(d2.day))
        else:
            breakdown.append(str(month_days[m]))
    return d1, d2, "＋".join(breakdown)

# ---------------------------------------------------------
# 3. 全20パターン作問ロジック
# ---------------------------------------------------------
modes_list = ["切り上げ", "四捨五入", "切り捨て"]

def gen_p1(idx):
    mode = modes_list[idx % 3]
    if idx % 2 == 0:
        unit_base = Decimal(random.randint(2500, 5500)) / Decimal('100')
        gross_price = unit_base * Decimal('12') + Decimal(random.choice(['0.10', '0.20', '0.30', '0.50', '0.70']))
        qty_doz = random.randint(35, 85) * 10 + random.choice([2, 5, 7])
        profit_pct = random.randint(12, 26)
        rate = Decimal(random.randint(14000, 15800)) / Decimal('100')
        raw_val = (gross_price / Decimal('12')) * Decimal(qty_doz) * (Decimal('1') + Decimal(profit_pct)/Decimal('100')) * rate
        ans = round_yen(raw_val, mode)
        q_text = f"1グロス＄{gross_price:.2f}の商品を{qty_doz:,}ダース販売するため、原価の{profit_pct}％の利益を見込んで定価をつけた。定価は円貨でいくらですか。ただし、＄1＝￥{rate:.2f}とする。（計算の最終で円未満{mode}）"
        f_text = f"＄{gross_price:.2f}÷12×{qty_doz:,}ダース×(1＋{profit_pct/100:.2f})×￥{rate:.2f}＝￥{raw_val:,.3f} → ￥{ans:,}（円未満{mode}）"
    else:
        doz_price = Decimal(random.randint(3000, 48000)) / Decimal('100')
        qty_pcs = random.randint(120, 580) * 12
        profit_pct = random.randint(12, 25)
        rate = Decimal(random.randint(14000, 15800)) / Decimal('100')
        raw_val = (doz_price / Decimal('12')) * Decimal(qty_pcs) * (Decimal('1') + Decimal(profit_pct)/Decimal('100')) * rate
        ans = round_yen(raw_val, mode)
        q_text = f"1ダース＄{doz_price:.2f}の商品を{qty_pcs:,}個販売するため、原価の{profit_pct}％の利益を見込んで定価をつけた。定価は円貨でいくらですか。ただし、＄1＝￥{rate:.2f}とする。（計算の最終で円未満{mode}）"
        f_text = f"＄{doz_price:.2f}÷12×{qty_pcs:,}個×(1＋{profit_pct/100:.2f})×￥{rate:.2f}＝￥{raw_val:,.3f} → ￥{ans:,}（円未満{mode}）"
    return "1.定価算出", "①外貨・グロス/ダース換算", q_text, f"￥{ans:,}", f_text

def gen_p2(idx):
    mode = modes_list[idx % 3]
    total_kg = random.randint(215, 785) * 10 + random.choice([5, 0])
    tare_pct = random.randint(2, 5)
    loss_pct = random.randint(1, 3)
    unit_price = random.randint(620, 980)
    profit_pct = random.randint(21, 29)
    net_kg = Decimal(total_kg) * (Decimal('1') - Decimal(tare_pct)/Decimal('100')) * (Decimal('1') - Decimal(loss_pct)/Decimal('100'))
    raw_val = Decimal(unit_price) * net_kg * (Decimal('1') + Decimal(profit_pct)/Decimal('100'))
    ans = round_yen(raw_val, mode)
    q_text = f"総量{total_kg:,}kg、風袋{tare_pct}％、減量は風袋を除いた量の{loss_pct}％の商品を、純量1kgにつき￥{unit_price:,}で仕入れ、仕入原価の{profit_pct}％の利益を見込んで定価をつけた。定価はいくらですか。（計算の最終で円未満{mode}）"
    f_text = f"{total_kg:,}kg×(1－{tare_pct/100:.2f})×(1－{loss_pct/100:.2f})＝純量{net_kg}kg / ￥{unit_price:,}×{net_kg}kg×(1＋{profit_pct/100:.2f})＝￥{raw_val:,.3f} → ￥{ans:,}（円未満{mode}）"
    return "1.定価算出", "②風袋・減量からの純量計算", q_text, f"￥{ans:,}", f_text

def gen_p3(idx):
    total_cost = random.randint(450, 980) * 1000
    expenses = random.randint(15, 85) * 1000
    cost = total_cost - expenses
    rate_wari = random.randint(1, 3)
    rate_bu = random.randint(1, 9)
    profit_rate = Decimal(rate_wari * 10 + rate_bu) / Decimal('100')
    selling_price = int(Decimal(total_cost) * (Decimal('1') + profit_rate))
    q_text = f"原価￥{cost:,}の商品を仕入れ、仕入諸掛￥{expenses:,}を支払った。この商品を￥{selling_price:,}で販売すると、利益率は諸掛込原価の何割何分ですか。"
    f_text = f"｛￥{selling_price:,}－(￥{cost:,}＋￥{expenses:,})｝÷(￥{cost:,}＋￥{expenses:,})＝{profit_rate:.2f}＝{rate_wari}割{rate_bu}分"
    return "2.損益・売買", "③諸掛込原価の利益率", q_text, f"{rate_wari}割{rate_bu}分", f_text

def gen_p4(idx):
    base_kg = random.choice([1, 10, 15])
    inc_pct = random.choice([2, 4, 5, 8])
    disc_pct = inc_pct - random.choice([1, -1])
    if disc_pct <= 1: disc_pct = 2
    unit_inc_base = random.randint(300, 900) * 100
    base_price = unit_inc_base * base_kg * (100 + inc_pct) // 100
    price_disc_per_kg = base_price * (100 - disc_pct) // (100 * base_kg)
    price_inc_per_kg = (base_price * 100) // (base_kg * (100 + inc_pct))
    diff = abs(price_disc_per_kg - price_inc_per_kg)
    winner = "値引" if price_disc_per_kg > price_inc_per_kg else "増量"
    q_text = f"{base_kg}kgにつき￥{base_price:,}の商品を販売するとき、値段を{disc_pct}％引きするのと数量を{inc_pct}％増量するのとでは、1kgあたりで比較すると、売り手にとってどちらがいくら有利ですか。（値引または増量を○で囲むこと）"
    f_text = f"値引時1kgあたり￥{price_disc_per_kg:,}、増量時1kgあたり￥{price_inc_per_kg:,} → 差額￥{diff:,}（{winner}有利）"
    return "2.損益・売買", "④値引と増量の有利比較", q_text, f"（{winner}）の方が ￥{diff:,} 有利", f_text

def gen_p5(idx):
    cost = random.randint(90, 850) * 10000
    markup = random.randint(35, 95) * 10000
    list_price = cost + markup
    wari = random.randint(1, 3)
    bu = random.randint(1, 8)
    disc_rate = Decimal(wari * 10 + bu) / Decimal('100')
    selling_price = int(Decimal(list_price) * (Decimal('1') - disc_rate))
    diff = selling_price - cost
    if diff == 0:
        markup += 100000
        list_price = cost + markup
        diff = int(Decimal(list_price) * (Decimal('1') - disc_rate)) - cost
    is_profit = diff > 0
    pl_str = f"￥{abs(diff):,}の利益" if is_profit else f"￥{abs(diff):,}の損失"
    ask_target = "定価" if idx % 2 == 0 else "原価"
    ans = list_price if ask_target == "定価" else cost
    q_text = f"原価に￥{markup:,}の利益を見込んで定価をつけ、定価の{wari}割{bu}分引きで販売したところ{pl_str}になった。この商品の{ask_target}はいくらですか。"
    f_text = f"原価をXとすると、(X＋￥{markup:,})×(1－{disc_rate:.2f})－X＝{'￥' if is_profit else '△￥'}{abs(diff):,} → 原価￥{cost:,}、定価￥{list_price:,}"
    return "2.損益・売買", "⑤損益からの定価・原価逆算", q_text, f"￥{ans:,}", f_text

def gen_p6(idx):
    disc_pct = random.choice([4, 5, 8, 10, 12, 15, 20])
    profit_pct = random.randint(12, 28)
    base_factor = random.randint(40, 120) * 100
    cost = base_factor * (100 - disc_pct)
    list_price = base_factor * (100 + profit_pct)
    q_text = f"仕入原価￥{cost:,}の商品を定価の{disc_pct}％引きで売っても、なお仕入原価の{profit_pct}％の利益を得たい。定価をいくらにすればよいですか。"
    f_text = f"定価をXとすると、X×(1－{disc_pct/100:.2f})＝￥{cost:,}×(1＋{profit_pct/100:.2f}) → 定価￥{list_price:,}"
    return "2.損益・売買", "⑥目標利益からの定価逆算", q_text, f"￥{list_price:,}", f_text

def gen_p7(idx):
    trade_val = random.randint(120, 950) * 10000
    s_rate = Decimal(random.randint(18, 58)) / Decimal('10')
    b_rate = Decimal(random.randint(15, 55)) / Decimal('10')
    seller_net = int(Decimal(trade_val) * (Decimal('1') - s_rate / Decimal('100')))
    buyer_total = int(Decimal(trade_val) * (Decimal('1') + b_rate / Decimal('100')))
    q_text = f"仲立人が売り主から{s_rate}％、買い主から{b_rate}％の手数料を受け取る約束で商品売買の仲介をした。売り主の手取金が￥{seller_net:,}のとき、買い主の支払総額はいくらですか。"
    f_text = f"￥{seller_net:,}÷(1－{s_rate/100:.3f})×(1＋{b_rate/100:.3f})＝￥{buyer_total:,}"
    return "3.仲立人", "⑦手取金→支払総額", q_text, f"￥{buyer_total:,}", f_text

def gen_p8(idx):
    trade_val = random.randint(150, 980) * 10000
    s_rate = Decimal(random.randint(18, 58)) / Decimal('10')
    b_rate = Decimal(random.randint(15, 55)) / Decimal('10')
    seller_net = int(Decimal(trade_val) * (Decimal('1') - s_rate / Decimal('100')))
    buyer_total = int(Decimal(trade_val) * (Decimal('1') + b_rate / Decimal('100')))
    q_text = f"仲立人が売り主から{s_rate}％、買い主から{b_rate}％の手数料を受け取る約束で商品売買の仲介をした。買い主の支払総額が￥{buyer_total:,}のとき、売り主の手取金はいくらですか。"
    f_text = f"￥{buyer_total:,}÷(1＋{b_rate/100:.3f})×(1－{s_rate/100:.3f})＝￥{seller_net:,}"
    return "3.仲立人", "⑧支払総額→手取金", q_text, f"￥{seller_net:,}", f_text

def gen_p9(idx):
    trade_val = random.randint(120, 950) * 10000
    s_rate = Decimal(random.randint(18, 62)) / Decimal('10')
    b_rate = Decimal(random.randint(12, 54)) / Decimal('10')
    fee_total = int(Decimal(trade_val) * (s_rate + b_rate) / Decimal('100'))
    if idx % 2 == 0:
        seller_net = int(Decimal(trade_val) * (Decimal('1') - s_rate / Decimal('100')))
        q_text = f"仲立人が売り主から{s_rate}％、買い主から{b_rate}％の手数料を受け取る約束で商品売買の仲介をした。売り主の手取金が￥{seller_net:,}のとき、仲立人の受け取る手数料合計はいくらですか。"
        f_text = f"￥{seller_net:,}÷(1－{s_rate/100:.3f})×({s_rate/100:.3f}＋{b_rate/100:.3f})＝￥{fee_total:,}"
    else:
        buyer_total = int(Decimal(trade_val) * (Decimal('1') + b_rate / Decimal('100')))
        q_text = f"仲立人が売り主から{s_rate}％、買い主から{b_rate}％の手数料を受け取る約束で商品売買の仲介をした。買い主の支払総額が￥{buyer_total:,}のとき、仲立人の受け取る手数料合計はいくらですか。"
        f_text = f"￥{buyer_total:,}÷(1＋{b_rate/100:.3f})×({s_rate/100:.3f}＋{b_rate/100:.3f})＝￥{fee_total:,}"
    return "3.仲立人", "⑨手数料合計の逆算", q_text, f"￥{fee_total:,}", f_text

def gen_p10(idx):
    qty_kg = random.randint(45, 680) * 10
    price_per_10kg = random.randint(35, 110) * 100
    expenses = random.randint(18, 95) * 1000
    fee_rate = Decimal(random.choice([20, 25, 30, 35, 40, 50])) / Decimal('10')
    sales = (qty_kg // 10) * price_per_10kg
    remittance = int(Decimal(sales) * (Decimal('1') - fee_rate / Decimal('100')) - Decimal(expenses))
    fee_str = f"{int(fee_rate)}" if fee_rate == int(fee_rate) else f"{fee_rate}"
    q_text = f"商品{qty_kg:,}kgの販売を委託され、10kgにつき￥{price_per_10kg:,}で販売し、販売諸掛￥{expenses:,}を立て替えた。販売手数料を売上高の{fee_str}％とすると、委託主への送金額はいくらですか。"
    f_text = f"({qty_kg:,}kg÷10kg×￥{price_per_10kg:,})×(1－{fee_rate/100:.3f})－￥{expenses:,}＝￥{remittance:,}"
    return "4.売買委託", "⑩販売委託の送金額", q_text, f"￥{remittance:,}", f_text

def gen_p11(idx):
    sales = random.randint(150, 480) * 10000
    fee_rate = Decimal(random.choice([25, 30, 35, 40, 45, 50])) / Decimal('10')
    expenses = random.randint(60, 350) * 100
    remittance = int(Decimal(sales) * (Decimal('1') - fee_rate / Decimal('100')) - Decimal(expenses))
    q_text = f"ある商品の販売を委託したところ、売上高の{fee_rate}％の手数料と、販売諸掛￥{expenses:,}を差し引いて￥{remittance:,}を送金してきた。売上高はいくらになりますか。"
    f_text = f"売上高をXとすると、X×(1－{fee_rate/100:.3f})－￥{expenses:,}＝￥{remittance:,} → 売上高￥{sales:,}"
    return "4.売買委託", "⑪販売委託の売上高逆算", q_text, f"￥{sales:,}", f_text

def gen_p12(idx):
    buy_price = random.randint(750, 2500) * 1000
    expenses = random.randint(20, 75) * 1000
    fee_pct = random.choice([3, 4, 5, 6, 7])
    total_cost = buy_price + expenses
    claim = int(Decimal(total_cost) * (Decimal('1') + Decimal(fee_pct) / Decimal('100')))
    q_text = f"買付の委託を受けた商品を￥{buy_price:,}で買い入れ、諸掛￥{expenses:,}を立て替えた。手数料を諸掛込原価の{fee_pct}％とすると、委託主への請求額はいくらですか。"
    f_text = f"(￥{buy_price:,}＋￥{expenses:,})×(1＋{fee_pct/100:.2f})＝￥{claim:,}"
    return "4.売買委託", "⑫買付委託の請求額", q_text, f"￥{claim:,}", f_text

def gen_p13(idx):
    usd = Decimal(random.randint(450000, 2900000)) / Decimal('100')
    rate_usd = Decimal(random.randint(14200, 15800)) / Decimal('100')
    rate_eur = rate_usd + Decimal(random.randint(1000, 2200)) / Decimal('100')
    raw_eur = usd * rate_usd / rate_eur
    ans_eur = round_cent(raw_eur)
    q_text = f"＄{usd:,.2f}は、€（ユーロ）でいくらですか。ただし、＄1＝￥{rate_usd:.2f}、€1＝￥{rate_eur:.2f}とする。（セント未満四捨五入）"
    f_text = f"＄{usd:,.2f}×￥{rate_usd:.2f}÷￥{rate_eur:.2f}＝€{raw_eur:,.4f} → €{ans_eur:,.2f}"
    return "5.為替換算", "⑬ドル(＄)→ユーロ(€)", q_text, f"€{ans_eur:,.2f}", f_text

def gen_p14(idx):
    eur = Decimal(random.randint(250000, 950000)) / Decimal('100')
    rate_usd = Decimal(random.randint(14200, 15800)) / Decimal('100')
    rate_eur = rate_usd + Decimal(random.randint(1000, 2200)) / Decimal('100')
    raw_usd = eur * rate_eur / rate_usd
    ans_usd = round_cent(raw_usd)
    q_text = f"€{eur:,.2f}は、米貨でいくらですか。ただし、€1＝￥{rate_eur:.2f}、＄1＝￥{rate_usd:.2f}とする。（セント未満四捨五入）"
    f_text = f"€{eur:,.2f}×￥{rate_eur:.2f}÷￥{rate_usd:.2f}＝＄{raw_usd:,.4f} → ＄{ans_usd:,.2f}"
    return "5.為替換算", "⑭ユーロ(€)→米貨(＄)", q_text, f"＄{ans_usd:,.2f}", f_text

def gen_p15(idx):
    principal = random.randint(60, 880) * 10000
    rate = Decimal(random.choice([18, 24, 30, 36, 42, 48, 54])) / Decimal('10')
    years = random.randint(1, 3)
    months = random.randint(1, 11)
    total_months = years * 12 + months
    interest = int(Decimal(principal) * (rate / Decimal('100')) * Decimal(total_months) / Decimal('12'))
    total_amt = principal + interest
    q_text = f"￥{principal:,}を年利率{rate}％で借り入れ、期日に元利合計￥{total_amt:,}を支払った。借入期間は、何年何か月ですか。"
    f_text = f"(￥{total_amt:,}－￥{principal:,})÷(￥{principal:,}×{rate/100:.3f})×12か月＝{total_months}か月＝{years}年{months}か月"
    return "6.単利利息", "⑮借入期間（年月）逆算", q_text, f"{years}年{months}か月", f_text

def gen_p16(idx):
    m_factor = random.randint(5, 25)
    principal = m_factor * 365000
    days = random.randint(15, 28) * 5
    rate = Decimal(random.randint(21, 58)) / Decimal('10')
    interest = int(Decimal(principal) * (rate / Decimal('100')) * Decimal(days) / Decimal('365'))
    total_amt = principal + interest
    sm = random.randint(3, 8)
    sd = random.randint(5, 20)
    d1, d2, b_str = get_date_pair(sm, sd, days, is_ryotan=False)
    q_text = f"元金￥{principal:,}を{d1.month}月{d1.day}日から{d2.month}月{d2.day}日まで貸し付け、期日に元利合計￥{total_amt:,}を受け取った。貸付利率は年何パーセントですか。（片落とし）"
    f_text = f"{d1.month}/{d1.day}〜{d2.month}/{d2.day}（片落とし）＝{days}日間 / 利息￥{interest:,}÷(￥{principal:,}×{days}日÷365日)＝{rate}％"
    return "6.単利利息", "⑯貸付利率逆算（片落とし）", q_text, f"{rate}％", f_text

def gen_p17(idx):
    if idx % 2 == 0:
        principal = random.randint(350, 980) * 1000
        rate = Decimal(random.randint(22, 54)) / Decimal('10')
        days = random.randint(65, 135)
        sm = random.randint(3, 8)
        sd = random.randint(5, 22)
        d1, d2, b_str = get_date_pair(sm, sd, days, is_ryotan=False)
        raw_total = Decimal(principal) + Decimal(principal) * (rate / Decimal('100')) * Decimal(days) / Decimal('365')
        ans = round_yen(raw_total, "四捨五入")
        q_text = f"￥{principal:,}を年利率{rate}％で{d1.month}月{d1.day}日から{d2.month}月{d2.day}日まで貸し付けた。期日に受け取る元利合計はいくらですか。（片落とし、円未満四捨五入）"
        f_text = f"{days}日間 / ￥{principal:,}＋利息＝￥{raw_total:,.3f} → ￥{ans:,}"
    else:
        principal = random.randint(45, 180) * 12000
        rate = Decimal(random.randint(21, 48)) / Decimal('10')
        years = random.randint(1, 3)
        months = random.randint(1, 11)
        total_months = years * 12 + months
        interest = int(Decimal(principal) * (rate / Decimal('100')) * Decimal(total_months) / Decimal('12'))
        ans = principal + interest
        q_text = f"￥{principal:,}を年利率{rate}％で{years}年{months}か月間貸し付けた。期日に受け取る元利合計はいくらですか。"
        f_text = f"￥{principal:,}＋利息￥{interest:,}＝￥{ans:,}"
    return "6.単利利息", "⑰元利合計算出", q_text, f"￥{ans:,}", f_text

def gen_p18(idx):
    principal = random.randint(250, 480) * 20000
    rate = Decimal(random.choice([18, 24, 25, 32, 36, 42])) / Decimal('10')
    years = random.randint(1, 3)
    months = random.choice([3, 6, 8, 9])
    total_months = years * 12 + months
    interest = int(Decimal(principal) * (rate / Decimal('100')) * Decimal(total_months) / Decimal('12'))
    total_amt = principal + interest
    q_text = f"年利率{rate}％で{years}年{months}か月間借り入れ、期日に元利合計￥{total_amt:,}を支払った。元金はいくらですか。"
    f_text = f"元金をXとすると、X×(1＋{rate/100:.3f}×{total_months}か月÷12か月)＝￥{total_amt:,} → 元金￥{principal:,}"
    return "6.単利利息", "⑱元金逆算（年月）", q_text, f"￥{principal:,}", f_text

def gen_p19(idx):
    m_factor = random.randint(5, 25)
    face_val = m_factor * 365000
    days = random.randint(45, 115)
    rate = Decimal(random.choice([25, 28, 30, 32, 35, 40, 45])) / Decimal('10')
    discount = int(Decimal(face_val) * (rate / Decimal('100')) * Decimal(days) / Decimal('365'))
    proceeds = face_val - discount
    q_text = f"額面￥{face_val:,}の手形を、割引率年{rate}％で割り引き、手取金￥{proceeds:,}を受け取った。割引日数は何日ですか。"
    f_text = f"割引料￥{discount:,}÷(￥{face_val:,}×{rate/100:.3f}÷365日)＝{days}日"
    return "7.手形割引", "⑲割引日数逆算", q_text, f"{days}日", f_text

def gen_p20(idx):
    sub_type = idx % 3
    if sub_type == 0:
        m_factor = random.randint(5, 25)
        face_val = m_factor * 365000
        days = random.randint(12, 24) * 5
        rate = Decimal(random.randint(24, 58)) / Decimal('10')
        discount = int(Decimal(face_val) * (rate / Decimal('100')) * Decimal(days) / Decimal('365'))
        sm = random.randint(4, 8)
        sd = random.randint(5, 22)
        d1, d2, b_str = get_date_pair(sm, sd, days, is_ryotan=True)
        q_text = f"{d2.month}月{d2.day}日満期、額面￥{face_val:,}の手形を、{d1.month}月{d1.day}日に割り引き、割引料￥{discount:,}を支払った。割引率は年何％ですか。（両端入れ）"
        f_text = f"{days}日間 / 割引料￥{discount:,}÷(￥{face_val:,}×{days}日÷365日)＝{rate}％"
        ans_str = f"{rate}％"
    elif sub_type == 1:
        m_factor = random.randint(5, 25)
        face_val = m_factor * 365000
        rate = Decimal(random.choice([25, 28, 32, 35, 40, 45])) / Decimal('10')
        days = random.randint(60, 95)
        discount = int(Decimal(face_val) * (rate / Decimal('100')) * Decimal(days) / Decimal('365'))
        proceeds = face_val - discount
        sm = random.randint(4, 8)
        sd = random.randint(10, 25)
        d1, d2, b_str = get_date_pair(sm, sd, days, is_ryotan=True)
        q_text = f"額面￥{face_val:,}の手形を、割引率年{rate}％で割り引き、手取金￥{proceeds:,}を受け取った。割引日を{d1.month}月{d1.day}日とすると、満期日は何月何日ですか。（両端入れ）"
        f_text = f"割引日数＝{days}日 / {d1.month}月{d1.day}日から両端入れで{days}日後 → {d2.month}月{d2.day}日"
        ans_str = f"{d2.month}月{d2.day}日"
    else:
        proceeds = random.randint(320, 850) * 10000
        days = random.choice([65, 75, 80, 85, 90, 100])
        rate = Decimal(random.randint(26, 52)) / Decimal('10')
        denom = Decimal('1') - (rate / Decimal('100')) * Decimal(days) / Decimal('365')
        raw_face = Decimal(proceeds) / denom
        ans_face = round_up_100yen(raw_face)
        q_text = f"手形を振り出し、これを割り引いて手取金￥{proceeds:,}を得たい。割引日数{days}日、割引率を年{rate}％とすると、手形金額をいくらにすればよいですか。（手形金額の￥100未満切り上げ）"
        f_text = f"手形金額をXとすると、X－X×{rate/100:.3f}×{days}日÷365日＝￥{proceeds:,} → ￥{ans_face:,}"
        ans_str = f"￥{ans_face:,}"
    return "7.手形割引", "⑳割引率・日付・額面逆算", q_text, ans_str, f_text

generators = [
    gen_p1, gen_p2, gen_p3, gen_p4, gen_p5, gen_p6, gen_p7, gen_p8, gen_p9, gen_p10,
    gen_p11, gen_p12, gen_p13, gen_p14, gen_p15, gen_p16, gen_p17, gen_p18, gen_p19, gen_p20
]

pattern_names = [
    "① 外貨・グロス/ダース換算", "② 風袋・減量からの純量計算", "③ 諸掛込原価の利益率",
    "④ 値引と増量の有利比較", "⑤ 損益からの定価・原価逆算", "⑥ 目標利益からの定価逆算",
    "⑦ 仲立人 手取金→支払総額", "⑧ 仲立人 支払総額→手取金", "⑨ 仲立人 手数料合計の逆算",
    "⑩ 販売委託の送金額", "⑪ 販売委託の売上高逆算", "⑫ 買付委託の請求額",
    "⑬ 為替 ドル(＄)→ユーロ(€)", "⑭ 為替 ユーロ(€)→米貨(＄)", "⑮ 借入期間（年月）逆算",
    "⑯ 貸付利率逆算（片落とし）", "⑰ 元利合計算出", "⑱ 元金逆算（年月）",
    "⑲ 割引日数逆算", "⑳ 割引率・日付・額面逆算"
]

# ---------------------------------------------------------
# 4. Wordファイル（BytesIO）生成
# ---------------------------------------------------------
def create_word_bytes(questions, mode_name, enable_trans, target_lang):
    doc = docx.Document()
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        
    doc.styles['Normal'].font.name = 'メイリオ'
    doc.styles['Normal'].font.size = Pt(10.5)

    title = doc.add_paragraph()
    r_t = title.add_run("計算実務能力検定試験 2級対策 ドリル")
    r_t.font.size = Pt(16)
    r_t.font.bold = True
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    trans_info = f" ［解説翻訳: {target_lang}］" if enable_trans else ""
    sub = doc.add_paragraph(f"商業計算 第1問（文章問題） 全 {len(questions)} 問 ［モード: {mode_name}］{trans_info}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(20)

    # 問題編
    doc.add_heading("【問題編】", level=1)
    current_cat = ""
    for q_no, cat, pat, q_text, ans, f_text in questions:
        if cat != current_cat and mode_name == "分野別":
            current_cat = cat
            h2 = doc.add_heading(f"■ {current_cat}", level=2)
            h2.paragraph_format.space_before = Pt(12)
            
        p_q = doc.add_paragraph()
        p_q.paragraph_format.left_indent = Inches(0.2)
        p_q.paragraph_format.first_line_indent = Inches(-0.2)
        p_q.add_run(f"({q_no}) ").bold = True
        p_q.add_run(q_text)

        p_ans = doc.add_paragraph()
        p_ans.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_ans.paragraph_format.space_after = Pt(14)
        p_ans.add_run("答 ＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿").bold = True

    doc.add_page_break()

    # 解答編
    doc.add_heading("【模範解答・解説編】", level=1)
    for q_no, cat, pat, q_text, ans, f_text in questions:
        p_sol = doc.add_paragraph()
        p_sol.add_run(f"({q_no}) 正解： {ans}\n").bold = True
        r_jp = p_sol.add_run(f"    【解説】 {f_text}\n")
        r_jp.font.color.rgb = RGBColor(60, 60, 60)
        
        if enable_trans and target_lang in TRANSLATION_DICT:
            tr_res = translate_explanation(f_text, target_lang)
            r_tr = p_sol.add_run(f"    【{target_lang}】 {tr_res}\n")
            r_tr.font.color.rgb = RGBColor(0, 102, 204)
            r_tr.font.size = Pt(9.5)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ---------------------------------------------------------
# 5. メイン画面UI
# ---------------------------------------------------------
st.title("📝 計算実務能力検定 2級 問題作成＆練習アプリ")

tab1, tab2 = st.tabs(["🖨️ プリント作成（Word / CSV ダウンロード）", "📱 画面で問題を解く（Webドリル）"])

# =========================================================
# TAB 1: Word/CSV出力
# =========================================================
with tab1:
    st.subheader("印刷・配布用プリントの自動生成")
    col1, col2 = st.columns([1, 1])

    with col1:
        mode_choice = st.radio(
            "出題形式を選択:",
            ["分野別（全20パターン網羅）", "模擬試験形式（本試験 8問1セット）", "パターン個別選択（苦手分野のみ）"]
        )
        
        if mode_choice == "分野別（全20パターン網羅）":
            count_val = st.number_input("各パターンの出題数:", min_value=1, max_value=50, value=2)
            st.caption(f"合計 {count_val * 20} 問が生成されます")
            selected_indices = list(range(20))
        elif mode_choice == "模擬試験形式（本試験 8問1セット）":
            count_val = st.number_input("模擬テストのセット数（回分）:", min_value=1, max_value=20, value=3)
            st.caption(f"合計 {count_val * 8} 問（{count_val}回分）が生成されます")
            selected_indices = []
        else:
            count_val = st.number_input("選択したパターンの各出題数:", min_value=1, max_value=50, value=2)
            st.write("出題したいパターンにチェックを入れてください:")
            selected_indices = []
            cols = st.columns(2)
            for i, p_name in enumerate(pattern_names):
                with cols[i % 2]:
                    if st.checkbox(p_name, value=(i < 4), key=f"chk_{i}"):
                        selected_indices.append(i)
            st.caption(f"合計 {count_val * len(selected_indices)} 問が生成されます")

    with col2:
        st.write("留学生向け解説翻訳の設定:")
        selected_lang = st.selectbox("解説に併記する言語:", LANG_OPTIONS, index=0)
        has_trans = (selected_lang != "なし（日本語のみ）")

        st.divider()
        generate_clicked = st.button("🚀 問題セットを生成する", type="primary", use_container_width=True)

    if generate_clicked:
        questions = []
        q_no = 1
        
        if mode_choice == "分野別（全20パターン網羅）":
            for gen in generators:
                for idx in range(count_val):
                    cat, pat, q_t, a_t, f_t = gen(idx)
                    questions.append((q_no, cat, pat, q_t, a_t, f_t))
                    q_no += 1
        elif mode_choice == "模擬試験形式（本試験 8問1セット）":
            for set_idx in range(count_val):
                chosen_gens = [
                    random.choice([gen_p1, gen_p2]),
                    random.choice([gen_p3, gen_p4, gen_p5, gen_p6]),
                    random.choice([gen_p7, gen_p8, gen_p9]),
                    random.choice([gen_p10, gen_p11, gen_p12]),
                    random.choice([gen_p13, gen_p14]),
                    random.choice([gen_p15, gen_p16, gen_p17, gen_p18]),
                    gen_p19,
                    gen_p20
                ]
                for p_idx, gen in enumerate(chosen_gens, start=1):
                    cat, pat, q_t, a_t, f_t = gen(set_idx * 8 + p_idx)
                    questions.append((q_no, f"第{set_idx+1}回 問{p_idx}", pat, q_t, a_t, f_t))
                    q_no += 1
        else:
            if not selected_indices:
                st.error("パターンが1つも選択されていません。")
            else:
                for idx in selected_indices:
                    gen = generators[idx]
                    for c in range(count_val):
                        cat, pat, q_t, a_t, f_t = gen(c)
                        questions.append((q_no, cat, pat, q_t, a_t, f_t))
                        q_no += 1

        if questions:
            st.session_state["generated_questions"] = questions
            st.session_state["has_trans"] = has_trans
            st.session_state["target_lang"] = selected_lang
            st.session_state["mode_name"] = mode_choice

    if "generated_questions" in st.session_state:
        qs = st.session_state["generated_questions"]
        st.success(f"全 {len(qs)} 問の生成が完了しました！下のボタンからダウンロードできます。")
        
        # Wordダウンロード
        word_buf = create_word_bytes(
            qs,
            st.session_state["mode_name"],
            st.session_state["has_trans"],
            st.session_state["target_lang"]
        )
        
        # CSVダウンロード
        csv_buf = io.StringIO()
        writer = csv.writer(csv_buf)
        header = ["番号", "分類/回次", "パターン", "問題文", "模範解答", "解説"]
        if st.session_state["has_trans"]:
            header.append(f"解説({st.session_state['target_lang']})")
        writer.writerow(header)
        for r in qs:
            row_data = list(r)
            if st.session_state["has_trans"]:
                row_data.append(translate_explanation(r[5], st.session_state["target_lang"]))
            writer.writerow(row_data)

        dl_col1, dl_col2 = st.columns(2)
        with dl_col1:
            st.download_button(
                label="📥 Word文書（.docx）をダウンロード",
                data=word_buf,
                file_name=f"keisan2kyu_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary",
                use_container_width=True
            )
        with dl_col2:
            st.download_button(
                label="📥 CSVファイル（.csv）をダウンロード",
                data=csv_buf.getvalue().encode("utf-8-sig"),
                file_name=f"keisan2kyu_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                mime="text/csv",
                use_container_width=True
            )

        with st.expander("作成された問題のプレビュー（先頭3問）", expanded=True):
            for q_n, cat, pat, q_t, a_t, f_t in qs[:3]:
                st.markdown(f"**({q_n}) [{cat}]** {q_t}")
                st.markdown(f"**答:** `{a_t}`")
                st.caption(f"解説: {f_t}")
                if st.session_state["has_trans"]:
                    st.caption(f"翻訳: {translate_explanation(f_t, st.session_state['target_lang'])}")
                st.write("---")

# =========================================================
# TAB 2: オンラインWebドリル
# =========================================================
with tab2:
    st.subheader("電卓を叩いてその場で解く（Webドリル）")

    col_drill1, col_drill2 = st.columns([1, 1])
    with col_drill1:
        pattern_select = st.selectbox(
            "練習したいパターン:",
            ["🎲 全20パターンからランダム出題"] + pattern_names
        )
    with col_drill2:
        drill_lang = st.selectbox("解説の表示言語:", LANG_OPTIONS, index=1, key="drill_lang")

    # 新しい問題の読み込み
    if "current_q" not in st.session_state or st.button("🔄 別の問題を出す", use_container_width=True):
        if pattern_select == "🎲 全20パターンからランダム出題":
            gen_func = random.choice(generators)
        else:
            p_idx = pattern_names.index(pattern_select)
            gen_func = generators[p_idx]
        
        cat, pat, q_t, a_t, f_t = gen_func(random.randint(1, 9999))
        st.session_state["current_q"] = (cat, pat, q_t, a_t, f_t)
        st.session_state["answered"] = False

    cat, pat, q_t, a_t, f_t = st.session_state["current_q"]

    st.info(f"**【{cat} - {pat}】**\n\n{q_t}")

    # 解答入力
    user_ans = st.text_input("あなたの解答（数値のみ入力でOK）:", key="ans_input")
    
    if st.button("判定する", type="primary"):
        st.session_state["answered"] = True
        
        # 簡易正誤比較（記号やカンマを除去）
        clean_user = user_ans.replace(",", "").replace("￥", "").replace("€", "").replace("＄", "").replace("円", "").strip()
        clean_ans = a_t.replace(",", "").replace("￥", "").replace("€", "").replace("＄", "").replace("円", "").strip()

        if clean_user == clean_ans or user_ans.strip() == a_t.strip():
            st.success(f"⭕ 正解です！ 正解：{a_t}")
        else:
            st.error(f"❌ 不正解... 正解：{a_t}")

        st.markdown(f"**【計算式・解説】**\n\n{f_t}")
        if drill_lang != "なし（日本語のみ）":
            st.markdown(f"**【{drill_lang} 解説】**\n\n{translate_explanation(f_t, drill_lang)}")